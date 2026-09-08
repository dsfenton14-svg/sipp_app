from datetime import datetime
from tkinter import filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import datos_Sipp as db


def _formatear_fecha(valor):
    if hasattr(valor, "strftime"):
        return valor.strftime("%d/%m/%Y")

    texto = str(valor or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato).strftime("%d/%m/%Y")
        except ValueError:
            continue
    return texto


def _parsear_hora(valor):
    texto = str(valor or "").strip().upper()
    if not texto:
        return None

    for formato in ("%I:%M %p", "%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def _calcular_horas(entrada_txt, salida_txt):
    inicio = _parsear_hora(entrada_txt)
    fin = _parsear_hora(salida_txt)
    if not inicio or not fin:
        return 0.0, 0.0

    if fin < inicio:
        from datetime import timedelta
        fin = fin + timedelta(days=1)

    horas = max(0.0, (fin - inicio).total_seconds() / 3600)
    extras = max(0.0, horas - 8.0)
    return horas, extras


def _estilizar_hoja(hoja, total_columnas):
    borde_suave = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    encabezado_fill = PatternFill("solid", fgColor="0F766E")
    encabezado_font = Font(color="FFFFFF", bold=True)
    titulo_font = Font(size=16, bold=True, color="0F172A")
    subtitulo_font = Font(size=11, italic=True, color="475569")
    ultima_columna = get_column_letter(total_columnas)

    hoja.merge_cells(f"A1:{ultima_columna}1")
    hoja["A1"] = "Reporte de Corte de Planilla"
    hoja["A1"].font = titulo_font
    hoja["A1"].alignment = Alignment(horizontal="center", vertical="center")

    hoja.merge_cells(f"A2:{ultima_columna}2")
    hoja["A2"].font = subtitulo_font
    hoja["A2"].alignment = Alignment(horizontal="center", vertical="center")

    for celda in hoja[4]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde_suave

    for fila in hoja.iter_rows(min_row=5, max_row=hoja.max_row, min_col=1, max_col=total_columnas):
        for celda in fila:
            celda.border = borde_suave
            celda.alignment = Alignment(horizontal="left", vertical="center")

    anchos = (16, 34, 18, 16, 16, 18, 16)
    for indice, ancho in enumerate(anchos[:total_columnas], start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    hoja.freeze_panes = "A5"
    hoja.auto_filter.ref = f"A4:{ultima_columna}{max(hoja.max_row, 4)}"


def _crear_hoja_resumen(libro, corte_id, fecha_inicio, fecha_fin, estado, resumen_empleados):
    hoja = libro.create_sheet("Resumen")
    hoja.append(["Resumen de horas por empleado"])
    hoja.append([f"Corte {corte_id} | Periodo: {_formatear_fecha(fecha_inicio)} a {_formatear_fecha(fecha_fin)} | Estado: {estado}"])
    hoja.append([f"Total de empleados: {len(resumen_empleados)}"])
    hoja.append(["Empleado", "DUI", "Marcaciones", "Horas trabajadas", "Horas extras"])

    for clave, datos in sorted(resumen_empleados.items(), key=lambda item: item[1]["nombre"].lower()):
        hoja.append([
            datos["nombre"],
            datos["dui"],
            datos["marcaciones"],
            round(datos["horas"], 2),
            round(datos["extras"], 2),
        ])

    _estilizar_hoja(hoja, 5)
    for celda in hoja.iter_rows(min_row=5, max_row=hoja.max_row, min_col=4, max_col=5):
        for item in celda:
            item.number_format = '0.00'


def descargar_corte_planilla_excel(corte):
    corte_id, fecha_inicio, fecha_fin, estado, *_ = corte
    if str(estado).upper() == "ACTIVO":
        messagebox.showwarning(
            "Corte activo",
            "Solo se puede descargar un corte que ya fue realizado y no está activo.",
        )
        return False

    archivo = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Workbook", "*.xlsx")],
        title=f"Guardar Corte {corte_id} como Excel",
        initialfile=f"corte_{corte_id}.xlsx",
    )
    if not archivo:
        return False

    servicio_marcaciones = db.marcaciones()
    marcaciones = servicio_marcaciones.obtener_marcaciones_por_rango(fecha_inicio, fecha_fin)

    libro = Workbook()
    hoja = libro.active
    hoja.title = f"Corte {corte_id}"
    hoja["A2"] = f"Corte {corte_id} | Periodo: {_formatear_fecha(fecha_inicio)} a {_formatear_fecha(fecha_fin)} | Estado: {estado}"
    hoja.append([])
    hoja.append(["Fecha", "Nombre", "DUI", "Entrada", "Salida", "Horas trabajadas", "Horas extras"])

    total_registros = 0
    resumen_empleados = {}
    for registro in marcaciones:
        if len(registro) < 6:
            continue
        _, nombre, dui, entrada, salida, fecha_registro = registro[:6]
        horas_trabajadas, horas_extras = _calcular_horas(entrada, salida)
        hoja.append([
            _formatear_fecha(fecha_registro),
            str(nombre or ""),
            str(dui or ""),
            str(entrada or ""),
            str(salida or ""),
            round(horas_trabajadas, 2),
            round(horas_extras, 2),
        ])
        clave = str(dui or "").strip() or str(nombre or "").strip()
        if clave not in resumen_empleados:
            resumen_empleados[clave] = {
                "nombre": str(nombre or ""),
                "dui": str(dui or ""),
                "marcaciones": 0,
                "horas": 0.0,
                "extras": 0.0,
            }
        resumen_empleados[clave]["marcaciones"] += 1
        resumen_empleados[clave]["horas"] += horas_trabajadas
        resumen_empleados[clave]["extras"] += horas_extras
        total_registros += 1

    _estilizar_hoja(hoja, 7)
    hoja["A3"] = f"Total de marcaciones: {total_registros}"
    hoja["A3"].font = Font(size=11, bold=True, color="0F766E")
    for fila in hoja.iter_rows(min_row=5, max_row=hoja.max_row, min_col=6, max_col=7):
        for celda in fila:
            celda.number_format = '0.00'

    _crear_hoja_resumen(libro, corte_id, fecha_inicio, fecha_fin, estado, resumen_empleados)

    libro.save(archivo)
    messagebox.showinfo(
        "Descarga completada",
        f"El corte {corte_id} se guardó correctamente en Excel.",
    )
    return True